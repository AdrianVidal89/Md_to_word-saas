"""Descubre/resuelve/prepara plantillas .docx de referencia.

Una plantilla es un .docx completo (portada, página de control de documento,
TOC, estilos de marca) que se usa como documento base en vez de crear uno en
blanco. Ver CLAUDE.md para la regla de seguridad sobre `resolve_template`.
"""

from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import styles


def list_templates() -> list:
    """Escanea templates/*.docx (excluye temporales de Word ~$*.docx)."""
    if not styles.TEMPLATES_DIR.exists():
        return []
    result = []
    for path in sorted(styles.TEMPLATES_DIR.glob("*.docx")):
        if path.name.startswith("~$"):
            continue
        result.append({"id": path.name, "name": path.stem})
    return result


def resolve_template(template_id: str):
    """Resuelve un ID de plantilla contra el almacén controlado (templates/).

    A diferencia del prototipo local original, esta función NUNCA acepta una
    ruta de filesystem arbitraria proveniente de un cliente: solo hace match
    contra un nombre de fichero dentro de TEMPLATES_DIR, y verifica que la
    ruta resuelta siga dentro de ese directorio. Cierra el riesgo de path
    traversal / arbitrary file read señalado en el análisis de arquitectura.
    """
    if not template_id or template_id == "none":
        return None

    candidate_name = Path(template_id).name  # descarta cualquier componente de ruta
    if not candidate_name.lower().endswith(".docx") or candidate_name.startswith("~$"):
        return None

    candidate = styles.TEMPLATES_DIR / candidate_name
    if not candidate.is_file():
        return None

    try:
        candidate.resolve().relative_to(styles.TEMPLATES_DIR.resolve())
    except ValueError:
        return None

    return candidate


def prepare_from_template(doc, metadata: dict) -> None:
    """Deja el documento base listo para recibir los bloques convertidos."""
    _strip_sample_content(doc)
    _fill_cover_fields(doc, metadata)
    _update_fields_on_open(doc)


def _strip_sample_content(doc) -> None:
    """Borra el contenido de ejemplo de la plantilla.

    Busca el primer párrafo con estilo TEMPLATE_CONTENT_MARKER_STYLE
    ("Heading 1") y, desde ahí (inclusive) hasta el final del body, borra
    todos los <w:p> y <w:tbl>, preservando siempre el <w:sectPr> final (no es
    ni w:p ni w:tbl, así que nunca se toca). Si no encuentra el marcador, no
    borra nada (modo degradado seguro).
    """
    marker_style = styles.TEMPLATE_CONTENT_MARKER_STYLE
    marker_element = None
    for paragraph in doc.paragraphs:
        if paragraph.style is not None and paragraph.style.name == marker_style:
            marker_element = paragraph._p
            break

    if marker_element is None:
        return

    body = doc.element.body
    remove = False
    for child in list(body):
        if child is marker_element:
            remove = True
        if remove and child.tag in (qn("w:p"), qn("w:tbl")):
            body.remove(child)


def _fill_cover_fields(doc, metadata: dict) -> None:
    """Rellena los campos de la página de control de documento.

    Para cada etiqueta conocida (COVER_FIELD_LABELS), sustituye el texto del
    siguiente párrafo no vacío tras la etiqueta, conservando el formato del
    primer run del párrafo destino.
    """
    labels = styles.COVER_FIELD_LABELS
    paragraphs = doc.paragraphs
    for idx, paragraph in enumerate(paragraphs):
        meta_key = labels.get(paragraph.text.strip())
        if not meta_key:
            continue
        value = metadata.get(meta_key)
        if value in (None, ""):
            continue
        for target in paragraphs[idx + 1:]:
            if target.text.strip():
                _set_paragraph_text(target, str(value))
                break


def _set_paragraph_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _update_fields_on_open(doc) -> None:
    """Fuerza a Word a recalcular campos automáticos (TOC incluido) al abrir."""
    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
